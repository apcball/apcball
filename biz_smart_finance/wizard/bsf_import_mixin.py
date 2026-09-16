"""CSV/XLSX parsing and live, server-checked previews for all finance imports."""
import base64
import binascii
import csv
import io
import math
from datetime import date, datetime
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from odoo import api, fields, models
from odoo.exceptions import UserError


class BsfImportMixin(models.AbstractModel):
    _name = "biz.smart.finance.import.mixin"
    _description = "Smart Finance — CSV/XLSX import helpers"

    _preview_fields = ()
    _preview_context_fields = ()

    company_id = fields.Many2one(
        "res.company", string="บริษัท", required=True, default=lambda self: self.env.company)
    currency_id = fields.Many2one(related="company_id.currency_id", readonly=True)
    file_data = fields.Binary(string="ไฟล์ (CSV / XLSX)")
    file_name = fields.Char(string="ชื่อไฟล์")
    state = fields.Selection([("draft", "อัปโหลด"), ("preview", "ตรวจสอบ")], default="draft")
    preview_result = fields.Json(compute="_compute_preview_result")
    valid_count = fields.Integer(compute="_compute_summary", string="แถวพร้อมนำเข้า")
    error_count = fields.Integer(compute="_compute_summary", string="แถวที่ต้องแก้")
    template_data = fields.Binary(readonly=True, attachment=False)
    template_name = fields.Char(readonly=True)

    def _file_columns(self):
        raise NotImplementedError

    def _required_columns(self):
        return ["amount"]

    def _read_rows(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError("กรุณาอัปโหลดไฟล์ก่อน")
        try:
            raw = base64.b64decode(self.file_data, validate=True)
        except (ValueError, binascii.Error):
            raise UserError("อ่านไฟล์ไม่ได้ กรุณาอัปโหลดไฟล์ใหม่")
        name = (self.file_name or "").lower()
        if name.endswith(".xlsx"):
            return self._read_xlsx(raw)
        if name.endswith(".csv") or "." not in name:
            return self._read_csv(raw)
        raise UserError("รองรับเฉพาะไฟล์ .csv หรือ .xlsx")

    @staticmethod
    def _cell_text(value):
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return str(value).strip() if value is not None else ""

    def _rows_from_cells(self, cells_rows):
        columns, required = self._file_columns(), self._required_columns()
        rows, header = [], []
        for idx, cells in enumerate(cells_rows, 1):
            values = [self._cell_text(c) for c in cells]
            if not any(values):
                continue
            if not header:
                # Empty trailing cells are ordinary Excel sheet formatting.
                while values and not values[-1]:
                    values.pop()
                header = [v.lower() for v in values]
                duplicates = {h for h in header if h and header.count(h) > 1}
                if duplicates:
                    raise UserError("แถว %s: หัวคอลัมน์ซ้ำ: %s" % (idx, ", ".join(sorted(duplicates))))
                if any(h not in columns for h in header) or any(c not in header for c in required):
                    raise UserError("แถว %s: หัวคอลัมน์ไม่ตรงรูปแบบ — รองรับ %s; จำเป็น: %s"
                                    % (idx, ", ".join(columns), ", ".join(required)))
                continue
            if any(values[len(header):]):
                raise UserError("แถว %s: มีข้อมูลเกินจำนวนหัวคอลัมน์" % idx)
            row = dict.fromkeys(columns, "")
            row.update(zip(header, values))
            row["row_index"] = idx
            rows.append(row)
        if not rows:
            raise UserError("ไม่พบแถวข้อมูลในไฟล์")
        return rows

    def _read_csv(self, raw):
        try:
            return self._rows_from_cells(csv.reader(io.StringIO(raw.decode("utf-8-sig")), strict=True))
        except UnicodeDecodeError:
            raise UserError("ไฟล์ CSV ต้องเป็น UTF-8")
        except csv.Error as error:
            raise UserError("รูปแบบ CSV ไม่ถูกต้อง: %s" % error)

    def _read_xlsx(self, raw):
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException
        except ImportError:
            raise UserError("อ่าน XLSX ต้องติดตั้ง openpyxl หรือบันทึกเป็น CSV แทน")
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            def cells():
                for row in workbook.active.iter_rows():
                    for cell in row:
                        if cell.data_type in ("f", "e"):
                            raise UserError("เซลล์ %s: กรุณาใช้ค่าข้อมูลแทนสูตรหรือค่า error ของ Excel" % cell.coordinate)
                    yield [cell.value for cell in row]
            return self._rows_from_cells(cells())
        except (BadZipFile, InvalidFileException, KeyError, ValueError, OSError, ParseError):
            raise UserError("ไฟล์ XLSX เสียหรือรูปแบบไม่ถูกต้อง กรุณาบันทึกไฟล์ใหม่")
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def _to_float(text):
        value = float(str(text).replace(",", "").strip() or 0.0)
        if not math.isfinite(value):
            raise ValueError("ต้องเป็นตัวเลขที่มีค่าจำกัด")
        return value

    @staticmethod
    def _to_date(text):
        return fields.Date.to_date(text.strip() or None)

    def _parse_cells(self, row, mappings):
        """Parse every cell independently; keep rejected input until that field is repaired."""
        values = {"row_index": row["row_index"]}
        errors = {}
        for target, (column, converter) in mappings.items():
            try:
                values[target] = converter(row[column])
            except (ValueError, TypeError):
                errors[target] = "%s: ค่า '%s' ไม่ถูกต้อง" % (column, row[column])
                values[target] = False
        values["input_errors"] = errors
        return values

    def _selection(self, value, allowed, default=False):
        value = (value or default or "").strip().lower()
        if value not in allowed:
            raise ValueError(value)
        return value

    @api.depends(lambda self: ["line_ids", "company_id"] +
                 ["line_ids." + name for name in (*self._preview_fields, "input_errors")] +
                 list(self._preview_context_fields))
    def _compute_preview_result(self):
        for wizard in self:
            context = wizard._validation_context()
            seen, results = set(), []
            for line in wizard.line_ids:
                values = {name: line[name] for name in wizard._preview_fields}
                error, key, is_new = wizard._row_error(values, context)
                # Failed numeric/date/selection conversions must never silently become zero/defaults.
                parse_errors = [message for name, message in (line.input_errors or {}).items()
                                if not line[name]]
                error = "; ".join(parse_errors) or error
                if key in seen:
                    error = error or "รายการซ้ำในตารางตัวอย่าง"
                if key is not None:
                    seen.add(key)
                results.append({"warning": error, "valid": not bool(error), "is_new_account": is_new and not error})
            wizard.preview_result = results

    def _validation_context(self):
        return {}

    @api.depends("preview_result")
    def _compute_summary(self):
        for wizard in self:
            rows = wizard.preview_result or []
            wizard.valid_count = sum(row["valid"] for row in rows)
            wizard.error_count = len(rows) - wizard.valid_count

    @api.onchange("file_data", "file_name")
    def _onchange_file(self):
        self.line_ids = [fields.Command.clear()]
        self.state = "draft"

    def write(self, vals):
        if {"file_data", "file_name", "mode"} & vals.keys():
            vals = dict(vals, line_ids=[fields.Command.clear()], state="draft")
        return super().write(vals)

    def action_preview(self):
        self.ensure_one()
        values = self._parse_rows(self._read_rows())
        self.write({"line_ids": [fields.Command.clear()] + [fields.Command.create(v) for v in values],
                    "state": "preview"})
        return self._reopen()

    def action_validate(self):
        self.ensure_one()
        self._compute_preview_result()
        return self._reopen()

    def action_import(self):
        self._check_import()
        with self.env.cr.savepoint():
            return self._import_records()

    def _check_import(self):
        self.ensure_one()
        if self.company_id not in self.env.companies:
            raise UserError("กรุณาเลือกบริษัทที่เปิดใช้งานและมีสิทธิ์เข้าถึง")
        if self.state != "preview" or not self.line_ids:
            raise UserError("ไม่มีแถวให้นำเข้า — กดตรวจไฟล์ก่อน")
        # Always recompute: related records may have changed since preview.
        self._compute_preview_result()
        errors = ["แถว %s: %s" % (line.row_index or index, result["warning"])
                  for index, (line, result) in enumerate(zip(self.line_ids, self.preview_result), 1)
                  if not result["valid"]]
        if errors:
            raise UserError("กรุณาแก้ข้อมูลก่อนนำเข้า\n" + "\n".join(errors[:20]))

    def _reopen(self):
        return {"type": "ir.actions.act_window", "res_model": self._name,
                "res_id": self.id, "view_mode": "form", "target": "new"}

    def action_template_csv(self):
        return self._download_template("csv")

    def action_template_xlsx(self):
        return self._download_template("xlsx")

    def _download_template(self, extension):
        self.ensure_one()
        from .bsf_import_templates import build_template
        content, name = build_template(self, extension)
        self.write({"template_data": base64.b64encode(content), "template_name": name})
        return {"type": "ir.actions.act_url", "target": "download",
                "url": "/web/content/%s/%s/template_data/%s?download=true" % (self._name, self.id, name)}


class BsfImportLineMixin(models.AbstractModel):
    _name = "biz.smart.finance.import.line.mixin"
    _description = "Smart Finance — live preview validation"

    input_errors = fields.Json(readonly=True)
    warning = fields.Char(string="ปัญหา", compute="_compute_validation")
    valid = fields.Boolean(compute="_compute_validation")

    @api.depends("wizard_id.preview_result")
    def _compute_validation(self):
        for wizard in self.mapped("wizard_id"):
            for line, result in zip(wizard.line_ids, wizard.preview_result or []):
                line.warning = result["warning"]
                line.valid = result["valid"]
        for line in self.filtered(lambda row: not row.wizard_id):
            line.warning = ""
            line.valid = False

    def write(self, vals):
        # A saved explicit edit (including correcting an invalid number to zero)
        # acknowledges only errors belonging to the edited fields.
        for line in self:
            errors = {name: message for name, message in (line.input_errors or {}).items() if name not in vals}
            super(BsfImportLineMixin, line).write(dict(vals, input_errors=errors))
        return True
