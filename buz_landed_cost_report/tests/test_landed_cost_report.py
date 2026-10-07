from odoo.tests import tagged
from odoo.tests.common import TransactionCase


@tagged('post_install', '-at_install')
class TestLandedCostReport(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.company = cls.env.company
        categ = cls.env['product.category'].create({
            'name': 'LC Test Categ',
            'property_cost_method': 'fifo',
            'property_valuation': 'manual_periodic',
        })
        cls.product = cls.env['product.product'].create({
            'name': 'LC Test Product',
            'type': 'product',
            'categ_id': categ.id,
            'standard_price': 0.0,
        })
        cls.service = cls.env['product.product'].create({
            'name': 'LC Test Freight',
            'type': 'service',
            'landed_cost_ok': True,
            'split_method_landed_cost': 'by_quantity',
        })
        wh = cls.env['stock.warehouse'].search([('company_id', '=', cls.company.id)], limit=1)
        cls.picking = cls.env['stock.picking'].create({
            'picking_type_id': wh.in_type_id.id,
            'location_id': cls.env.ref('stock.stock_location_suppliers').id,
            'location_dest_id': wh.lot_stock_id.id,
            'move_ids': [(0, 0, {
                'name': 'LC test move',
                'product_id': cls.product.id,
                'product_uom_qty': 10,
                'product_uom': cls.product.uom_id.id,
                'price_unit': 100.0,
                'location_id': cls.env.ref('stock.stock_location_suppliers').id,
                'location_dest_id': wh.lot_stock_id.id,
            })],
        })
        cls.picking.action_confirm()
        cls.picking.move_ids.quantity = 10
        cls.picking.move_ids.picked = True
        cls.picking._action_done()

    def _make_lc(self, amounts, types=None):
        journal = self.env['account.journal'].search(
            [('type', '=', 'general'), ('company_id', '=', self.company.id)], limit=1)
        expense = self.env['account.account'].search(
            [('account_type', '=', 'expense'), ('company_id', '=', self.company.id)], limit=1)
        lc = self.env['stock.landed.cost'].create({
            'account_journal_id': journal.id,
            'picking_ids': [(6, 0, self.picking.ids)],
            'cost_lines': [(0, 0, dict({
                'name': 'Cost %s' % i,
                'product_id': self.service.id,
                'price_unit': amt,
                'split_method': 'by_quantity',
                'account_id': expense.id,
            }, **({'cost_type_id': types[i % len(types)].id} if types else {})))
                for i, amt in enumerate(amounts)],
        })
        lc.compute_landed_cost()
        self.env.flush_all()  # SQL views read tables directly
        return lc

    def test_report_row_and_detail(self):
        lc = self._make_lc([50.0, 25.0])
        lc.button_validate()
        self.env.flush_all()
        rows = self.env['buz.landed.cost.report'].search([('landed_cost_id', '=', lc.id)])
        self.assertEqual(len(rows), 1)
        self.assertAlmostEqual(rows.qty, 10)
        self.assertAlmostEqual(rows.base_cost, 1000.0)
        self.assertAlmostEqual(rows.landed_cost, 75.0)
        self.assertAlmostEqual(rows.total_cost, 1075.0)
        self.assertAlmostEqual(rows.unit_cost, 107.5)
        self.assertAlmostEqual(sum(rows.detail_ids.mapped('amount')), 75.0)
        self.assertEqual(len(rows.detail_ids), 2)
        # no FX applied: landed amount equals what the LC allocated
        self.assertAlmostEqual(rows.svl_landed_value, 75.0)
        self.assertAlmostEqual(rows.svl_landed_diff, 0.0)

    def test_draft_flagged_and_audit_clean(self):
        draft = self._make_lc([10.0])
        row = self.env['buz.landed.cost.report'].search([('landed_cost_id', '=', draft.id)])
        self.assertEqual(row.state, 'draft')
        draft.unlink()
        lc = self._make_lc([40.0])
        lc.button_validate()
        self.env.flush_all()
        audit = self.env['buz.landed.cost.audit'].search([('landed_cost_id', '=', lc.id)])
        self.assertAlmostEqual(audit.amount_total, 40.0)
        self.assertAlmostEqual(audit.alloc_total, 40.0)
        self.assertAlmostEqual(audit.svl_total, 40.0)
        self.assertFalse(audit.alloc_mismatch)
        self.assertFalse(audit.svl_over)

    def test_wizard_default_excludes_draft(self):
        wiz = self.env['buz.landed.cost.report.wizard'].create({})
        self.assertIn(('state', '=', 'done'), wiz._get_domain())

    def test_consumed_before_lc_is_not_a_defect(self):
        """Odoo capitalises LC only for stock on hand: 4 of 10 sold -> SVL = 60%."""
        wh = self.picking.picking_type_id.warehouse_id
        out = self.env['stock.picking'].create({
            'picking_type_id': wh.out_type_id.id,
            'location_id': wh.lot_stock_id.id,
            'location_dest_id': self.env.ref('stock.stock_location_customers').id,
            'move_ids': [(0, 0, {
                'name': 'LC test out',
                'product_id': self.product.id,
                'product_uom_qty': 4,
                'product_uom': self.product.uom_id.id,
                'location_id': wh.lot_stock_id.id,
                'location_dest_id': self.env.ref('stock.stock_location_customers').id,
            })],
        })
        out.action_confirm()
        out.move_ids.quantity = 4
        out.move_ids.picked = True
        out._action_done()
        lc = self._make_lc([100.0])
        lc.button_validate()
        self.env.flush_all()
        row = self.env['buz.landed.cost.report'].search([('landed_cost_id', '=', lc.id)])
        self.assertAlmostEqual(row.landed_cost, 100.0)
        self.assertAlmostEqual(row.svl_landed_value, 60.0)
        self.assertAlmostEqual(row.svl_landed_diff, 40.0)
        audit = self.env['buz.landed.cost.audit'].search([('landed_cost_id', '=', lc.id)])
        self.assertAlmostEqual(audit.uncapitalised, 40.0)
        self.assertFalse(audit.svl_over)
        self.assertFalse(audit.alloc_mismatch)

    def test_cost_type_breakdown_and_custom_type(self):
        freight = self.env['buz.landed.cost.type'].create({'name': 'Freight Test', 'code': 'frt'})
        duty = self.env.ref('buz_landed_cost_report.cost_type_tax')
        lc = self._make_lc([30.0, 20.0], types=freight | duty)
        lc.button_validate()
        self.env.flush_all()
        det = self.env['buz.landed.cost.report.detail'].search([('landed_cost_id', '=', lc.id)])
        by_type = {d.cost_type_id.name: d.amount for d in det}
        self.assertAlmostEqual(by_type['Freight Test'], 30.0)
        self.assertAlmostEqual(by_type['Tax'], 20.0)

    def test_default_cost_type_is_expense(self):
        lc = self._make_lc([10.0])
        self.assertEqual(lc.cost_lines.cost_type_id,
                         self.env.ref('buz_landed_cost_report.cost_type_expense'))

    def test_cost_type_follows_product(self):
        freight = self.env['buz.landed.cost.type'].create({'name': 'Freight P', 'code': 'frtp'})
        self.service.product_tmpl_id.landed_cost_type_id = freight
        lc = self._make_lc([10.0])
        self.assertEqual(lc.cost_lines.cost_type_id, freight)
        lc.button_validate()
        self.env.flush_all()
        det = self.env['buz.landed.cost.report.detail'].search([('landed_cost_id', '=', lc.id)])
        self.assertEqual(det.cost_type_id, freight)

    def _export_xlsx(self, wizard):
        report = self.env.ref('buz_landed_cost_report.action_report_landed_cost_xlsx')
        content, _fmt = report._render_xlsx(
            report.report_name, wizard.ids,
            {'wizard_id': wizard.id, 'domain': wizard._get_domain()})
        return content

    def test_xlsx_export_styled(self):
        import io
        from openpyxl import load_workbook
        lc = self._make_lc([50.0, 25.0])
        lc.button_validate()
        self.env.flush_all()
        wiz = self.env['buz.landed.cost.report.wizard'].create({'landed_cost_ids': [(6, 0, lc.ids)]})
        wb = load_workbook(io.BytesIO(self._export_xlsx(wiz)))
        self.assertEqual(wb.sheetnames, ['Summary', 'Cost Breakdown', 'Audit'])
        ws = wb['Summary']
        self.assertEqual(ws['A6'].value, 'Landed Cost')  # header row
        self.assertEqual(ws['A6'].fill.fgColor.rgb[-6:], '1F4E78')
        self.assertEqual(ws.freeze_panes, 'C7')
        self.assertEqual(ws['A7'].value, lc.name)
        self.assertEqual(ws['C7'].value, 'Posted')
        self.assertEqual(ws['A8'].value, 'TOTAL')
        self.assertIn('SUBTOTAL', str(ws['I8'].value))

    def test_xlsx_export_empty(self):
        wiz = self.env['buz.landed.cost.report.wizard'].create({
            'date_from': '1990-01-01', 'date_to': '1990-01-02'})
        self.assertTrue(self._export_xlsx(wiz))
