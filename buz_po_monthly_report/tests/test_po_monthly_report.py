import io
import zipfile
from datetime import date

from odoo.exceptions import ValidationError
from odoo.tests import TransactionCase, tagged


@tagged('post_install', '-at_install')
class TestPurchaseOrderMonthlyReport(TransactionCase):

    def test_report_headers_match_reference_order(self):
        report = self.env['buz.po.monthly.report.xlsx.generator']
        self.assertEqual(report.HEADERS, [
            'ลำดับ', 'เลขที่เอกสาร', 'วันที่เปิดPO', 'รหัสผู้จำหน่าย',
            'ชื่อผู้จำหน่าย', 'Ref', 'กำหนดส่ง', 'วันที่ Approve', 'หมายเหตุ',
            'ชื่อ', 'ราคา', 'จำนวน', 'AMOUNT', 'รับ', 'คงเหลือ', 'สถานที่ส่ง',
            'CREDIT',
        ])

    def test_selected_pr_states_returns_only_checked_states(self):
        wizard = self.env['po.monthly.report.wizard'].new({
            'pr_state_waiting_purchase_approval': True,
            'pr_state_approved': True,
        })
        report = self.env['buz.po.monthly.report.xlsx.generator']
        self.assertEqual(
            report._selected_pr_states(wizard),
            ['waiting_purchase_approval', 'approved'],
        )

    def test_no_selected_pr_state_disables_status_filter(self):
        wizard = self.env['po.monthly.report.wizard'].new({})
        report = self.env['buz.po.monthly.report.xlsx.generator']
        self.assertEqual(report._selected_pr_states(wizard), [])

    def test_at_least_one_complete_date_range_is_required(self):
        wizard_model = self.env['po.monthly.report.wizard']
        with self.assertRaises(ValidationError):
            wizard_model.new({})._check_date_ranges()

        wizard = wizard_model.new({
            'po_date_from': date(2026, 5, 1),
            'po_date_to': date(2026, 5, 31),
        })
        wizard._check_date_ranges()

    def test_partial_and_reversed_ranges_are_rejected(self):
        wizard_model = self.env['po.monthly.report.wizard']
        partial = wizard_model.new({'po_date_from': date(2026, 5, 1)})
        with self.assertRaises(ValidationError):
            partial._check_date_ranges()

        reversed_range = wizard_model.new({
            'po_date_from': date(2026, 5, 31),
            'po_date_to': date(2026, 5, 1),
        })
        with self.assertRaises(ValidationError):
            reversed_range._check_date_ranges()


    def test_pr_date_fields_are_not_supported(self):
        wizard = self.env['po.monthly.report.wizard']
        self.assertNotIn('pr_date_from', wizard._fields)
        self.assertNotIn('pr_date_to', wizard._fields)

    def test_export_creates_valid_xlsx_for_empty_result(self):
        wizard = self.env['po.monthly.report.wizard'].new({
            'po_date_from': date(1900, 1, 1),
            'po_date_to': date(1900, 1, 2),
        })
        content = self.env[
            'buz.po.monthly.report.xlsx.generator'
        ].generate_xlsx(wizard)
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(content)))

    def test_export_rows_skip_display_lines_and_keep_order(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest('openpyxl not installed')
        partner = self.env['res.partner'].with_context(
            skip_partner_required_fields=True,
        ).create({'name': 'PO Report Vendor'})
        product = self.env['product.product'].create({'name': 'PO Report Item'})
        order = self.env['purchase.order'].create({
            'partner_id': partner.id,
            'date_order': '2026-05-10 03:00:00',
            'order_line': [
                (0, 0, {'display_type': 'line_section', 'name': 'Section'}),
                (0, 0, {
                    'product_id': product.id, 'name': 'Item A',
                    'product_qty': 2, 'price_unit': 10,
                }),
                (0, 0, {
                    'product_id': product.id, 'name': 'Item B',
                    'product_qty': 1.5, 'price_unit': 20,
                }),
            ],
        })
        wizard = self.env['po.monthly.report.wizard'].new({
            'po_date_from': date(2026, 5, 10),
            'po_date_to': date(2026, 5, 10),
        })
        content = self.env[
            'buz.po.monthly.report.xlsx.generator'
        ].generate_xlsx(wizard)
        sheet = load_workbook(io.BytesIO(content)).active
        rows = [r for r in sheet.iter_rows(min_row=5, values_only=True) if r[1] == order.name]
        self.assertEqual([r[9] for r in rows], ['Item A', 'Item B'])
        self.assertEqual([r[0] for r in rows][1] - [r[0] for r in rows][0], 1)
        self.assertEqual(rows[0][12], 20)
