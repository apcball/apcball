import base64
import csv
import io

from odoo import Command, _, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:
    openpyxl = None

REQUIRED_COLUMNS = ["product_code", "warehouse_code", "target_qty", "target_value"]


def _to_float(value):
    """Return float(value) or None when not numeric. Tolerates Excel-style
    thousands separators and surrounding whitespace."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


class StockCountAdjustmentImport(models.TransientModel):
    _name = "stock.count.adjustment.import"
    _description = "Stock Count Adjustment Import Wizard"

    adjustment_id = fields.Many2one(
        "stock.count.adjustment",
        required=True,
        default=lambda self: self.env.context.get("active_id"),
    )
    data_file = fields.Binary(string="File", required=True)
    filename = fields.Char()
    file_format = fields.Selection(
        [("xlsx", "Excel"), ("csv", "CSV")],
        required=True,
        default="xlsx",
    )
    import_valid_only = fields.Boolean(
        string="Import valid rows only",
        default=False,
        help="If unchecked, a single rejected row aborts the whole import.",
    )
    result_log = fields.Text(readonly=True)

    def _reopen(self):
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "view_mode": "form",
            "res_id": self.id,
            "target": "new",
        }

    def _read_rows(self):
        """Yield lists of cell values, first row = header."""
        raw = base64.b64decode(self.data_file)
        if self.file_format == "csv":
            reader = csv.reader(io.StringIO(raw.decode("utf-8-sig")))
            return [list(r) for r in reader]
        if not openpyxl:
            raise UserError(_("Python module 'openpyxl' is required for xlsx import."))
        try:
            wb = openpyxl.load_workbook(
                io.BytesIO(raw), read_only=True, data_only=True
            )
        except Exception as exc:  # noqa: BLE001
            raise UserError(_("Invalid Excel file: %s") % exc)
        ws = wb[wb.sheetnames[0]]
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def _resolve_product(self, code):
        products = self.env["product.product"].search(
            [("default_code", "=", code)]
        )
        if not products:
            return self.env["product.product"]
        scoped = products.filtered(
            lambda p: not p.company_id
            or p.company_id == self.adjustment_id.company_id
        )
        return (scoped or products)[:1]

    def _resolve_warehouse(self, code):
        code = code.split("/")[0].strip()
        whs = self.env["stock.warehouse"].search([("code", "=", code)])
        if not whs:
            return self.env["stock.warehouse"]
        scoped = whs.filtered(
            lambda w: w.company_id == self.adjustment_id.company_id
        )
        return (scoped or whs)[:1]

    def action_do_import(self):
        self.ensure_one()
        if self.adjustment_id.state != "draft":
            raise UserError(
                _("Lines can only be imported into a draft adjustment.")
            )

        rows = self._read_rows()
        if not rows:
            raise UserError(_("The file is empty."))

        header = [
            str(h).strip().lower() if h is not None else "" for h in rows[0]
        ]
        missing = [c for c in REQUIRED_COLUMNS if c not in header]
        if missing:
            raise UserError(
                _("Missing required column(s): %s") % ", ".join(missing)
            )
        idx = {name: header.index(name) for name in header if name}

        def cell(row, key):
            i = idx.get(key)
            if i is None or i >= len(row):
                return None
            return row[i]

        resolved = []  # list of dicts: product_id, warehouse_id, qty, value, note
        rejects = []

        for rownum, row in enumerate(rows[1:], start=2):
            if not row or all(v is None or v == "" for v in row):
                continue

            code = cell(row, "product_code")
            code = str(code).strip() if code is not None else ""
            wh_code = cell(row, "warehouse_code")
            wh_code = str(wh_code).strip() if wh_code is not None else ""
            qty = _to_float(cell(row, "target_qty"))
            value = _to_float(cell(row, "target_value"))
            note = cell(row, "note")
            note = str(note).strip() if note is not None else ""

            product = self._resolve_product(code) if code else \
                self.env["product.product"]
            warehouse = self._resolve_warehouse(wh_code) if wh_code else \
                self.env["stock.warehouse"]

            row_errs = []
            if not product:
                row_errs.append(_("unknown product_code '%s'") % code)
            if not warehouse:
                row_errs.append(_("unknown warehouse_code '%s'") % wh_code)
            if qty is None:
                row_errs.append(_("non-numeric target_qty"))
            if value is None:
                row_errs.append(_("non-numeric target_value"))

            if row_errs:
                rejects.append(_("Row %s: %s") % (rownum, "; ".join(row_errs)))
                continue

            resolved.append({
                "product_id": product.id,
                "warehouse_id": warehouse.id,
                "target_qty": qty,
                "target_value": value,
                "note": note,
            })

        # bucket_seq per (product, warehouse) group, in file order, seeded from
        # the max existing bucket_seq so a second import does not collide with
        # the unique(adjustment, product, warehouse, bucket_seq) constraint.
        seq_by_group = {}
        for line in self.adjustment_id.line_ids:
            key = (line.product_id.id, line.warehouse_id.id)
            seq_by_group[key] = max(seq_by_group.get(key, 0), line.bucket_seq)
        for item in resolved:
            key = (item["product_id"], item["warehouse_id"])
            seq_by_group[key] = seq_by_group.get(key, 0) + 10
            item["bucket_seq"] = seq_by_group[key]

        log_lines = []
        if rejects and not self.import_valid_only:
            log_lines.append(
                _("Import aborted: %s row(s) rejected, nothing imported.")
                % len(rejects)
            )
            log_lines.extend(rejects)
            self.write({"result_log": "\n".join(log_lines)})
            return self._reopen()

        if resolved:
            self.adjustment_id.write({
                "line_ids": [Command.create(item) for item in resolved]
            })

        log_lines.append(_("Imported %s line(s).") % len(resolved))
        if rejects:
            log_lines.append(_("Skipped %s row(s):") % len(rejects))
            log_lines.extend(rejects)
        self.write({"result_log": "\n".join(log_lines)})

        return {
            "type": "ir.actions.act_window",
            "res_model": "stock.count.adjustment",
            "res_id": self.adjustment_id.id,
            "view_mode": "form",
            "target": "current",
        }
