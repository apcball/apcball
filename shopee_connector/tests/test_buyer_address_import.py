import base64
import io

from odoo.exceptions import UserError
from odoo.tests import TransactionCase

try:
    import openpyxl
except ImportError:
    openpyxl = None

# Subset of the Seller Center "Order to ship" export, keeping the duplicated
# province/district/zip headers of the tax-invoice block.
HEADERS = [
    "หมายเลขคำสั่งซื้อ", "สถานะการสั่งซื้อ", "ชื่อผู้ใช้ (ผู้ซื้อ)", "ชื่อสินค้า",
    "ชื่อผู้รับ", "หมายเลขโทรศัพท์", "ที่อยู่ในการจัดส่ง", "ประเทศ", "จังหวัด",
    "เขต/อำเภอ", "รหัสไปรษณีย์", "เขต/อำเภอ", "จังหวัด", "รหัสไปรษณีย์",
    "อีเมลสำหรับรับใบกำกับภาษี",
]


class TestShopeeBuyerAddressImport(TransactionCase):
    def setUp(self):
        super().setUp()
        if not openpyxl:
            self.skipTest("openpyxl is not installed")
        self.config = self.env["shopee.config"].create({
            "name": "Test Shop", "environment": "sandbox",
            "partner_id": "1000001", "partner_key": "testkey", "shop_id": "222",
            "access_token": "tok", "token_expires_at": "2999-01-01 00:00:00",
        })
        self.order = self.env["sale.order"].create_from_shopee({
            "order_sn": "260921KGN5CTYC",
            "buyer_user_id": 777, "buyer_username": "465twj8odj",
            "region": "TH", "order_status": "READY_TO_SHIP",
            "recipient_address": {"name": "****", "phone": "****",
                                  "full_address": "****", "region": "****"},
            "item_list": [{"item_name": "Widget", "model_sku": "NO-SUCH-SKU"}],
        }, config=self.config)

    def _xlsx(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        return base64.b64encode(stream.getvalue())

    def _row(self, order_sn, product):
        return [
            order_sn, "ที่ต้องจัดส่ง", "465twj8odj", product,
            "เค้ก", "66644976555",
            "เลขที่ 63/50 หมู่ที่ 5 ตำบลบ่อผุด อำเภอเกาะสมุย จังหวัดสุราษฎร์ธานี 84320",
            "TH", "จังหวัดสุราษฎร์ธานี", "อำเภอเกาะสมุย", "84320",
            "TAX DISTRICT", "TAX PROVINCE", "99999", "",
        ]

    def test_import_fills_buyer_without_creating_orders(self):
        order_count = self.env["sale.order"].search_count([])
        wizard = self.env["shopee.buyer.address.import.wizard"].create({
            "file_name": "Order.toship.xlsx",
            "file_data": self._xlsx([
                self._row("260921KGN5CTYC", "Item A"),
                self._row("260921KGN5CTYC", "Item B"),
                self._row("NOT-IN-ODOO-1", "Item C"),
            ]),
        })
        wizard.action_import()

        partner = self.order.partner_id
        self.assertEqual(partner.name, "เค้ก")
        self.assertEqual(partner.phone, "0644976555")
        # trailing ตำบล/อำเภอ/จังหวัด/zip are split out, not repeated
        self.assertEqual(partner.street, "เลขที่ 63/50 หมู่ที่ 5")
        self.assertEqual(partner.street2, "ตำบลบ่อผุด")
        self.assertEqual(partner.city, "อำเภอเกาะสมุย")
        self.assertEqual(partner.zip, "84320")  # shipping zip, not the tax one
        self.assertEqual(partner.country_id, self.env.ref("base.th"))
        self.assertEqual(self.env["sale.order"].search_count([]), order_count)
        self.assertEqual(wizard.state, "done")
        self.assertIn("1 order(s)", wizard.result_message)
        self.assertIn("NOT-IN-ODOO-1", wizard.result_message)

    def test_rejects_non_seller_center_file(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["SKU", "Stock"])
        stream = io.BytesIO()
        workbook.save(stream)
        wizard = self.env["shopee.buyer.address.import.wizard"].create({
            "file_name": "stock.xlsx",
            "file_data": base64.b64encode(stream.getvalue()),
        })
        with self.assertRaises(UserError):
            wizard.action_import()


TAX_HEADERS = HEADERS[:-1] + [
    "ผู้ซื้อร้องขอใบกำกับภาษี", "ประเภทใบกำกับภาษี", "ชื่อ", "ประเภทสาขา",
    "ชื่อสาขา", "รหัสประจำสาขา", "ที่อยู่สำหรับออกใบกำกับภาษีแบบเต็มรูป",
    "รายละเอียดที่อยู่", "แขวง/ตำบล", "หมายเลขประจำตัวผู้เสียภาษี",
    "หมายเลขโทรศัพท์สำหรับออกใบกำกับภาษี", "อีเมลสำหรับรับใบกำกับภาษี",
]


class TestShopeeBuyerTaxInvoiceImport(TestShopeeBuyerAddressImport):
    def _xlsx(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(TAX_HEADERS)
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        return base64.b64encode(stream.getvalue())

    def _row(self, order_sn, product, tax=True):
        # HEADERS[:-1] ends with the tax-invoice district/province/zip set
        base = [
            order_sn, "ที่ต้องจัดส่ง", "465twj8odj", product,
            "เค้ก", "66644976555",
            "เลขที่ 63/50 หมู่ที่ 5 ตำบลบ่อผุด อำเภอเกาะสมุย จังหวัดสุราษฎร์ธานี 84320",
            "TH", "จังหวัดสุราษฎร์ธานี", "อำเภอเกาะสมุย", "84320",
            "อำเภอชะอำ", "จังหวัดเพชรบุรี", "76120",
        ]
        if not tax:
            return base + ["No"] + [""] * 11
        return base + [
            "Yes", "Personal", "ธนวรรษ ดำรงค์แสง", "สาขาย่อย", "", "",
            "164/4 ม.6 ตำบลสามพระยา อำเภอชะอำ จังหวัดเพชรบุรี 76120",
            "164/4 ม.6", "ตำบลสามพระยา", "1234567890123", "",
            "buyer@example.com",
        ]

    def test_import_fills_buyer_without_creating_orders(self):
        """Without a tax invoice request the buyer gets the recipient data."""
        wizard = self.env["shopee.buyer.address.import.wizard"].create({
            "file_name": "Order.toship.xlsx",
            "file_data": self._xlsx([self._row("260921KGN5CTYC", "A", tax=False)]),
        })
        wizard.action_import()
        partner = self.order.partner_id
        self.assertEqual(partner.name, "เค้ก")
        self.assertEqual(partner.phone, "0644976555")
        self.assertFalse(partner.vat)
        self.assertEqual(self.order.partner_shipping_id, partner)

    def test_tax_invoice_splits_invoice_and_delivery(self):
        wizard = self.env["shopee.buyer.address.import.wizard"].create({
            "file_name": "Order.toship.xlsx",
            "file_data": self._xlsx([self._row("260921KGN5CTYC", "A")]),
        })
        wizard.action_import()

        customer = self.order.partner_id
        self.assertEqual(customer.name, "ธนวรรษ ดำรงค์แสง")
        self.assertEqual(customer.vat, "1234567890123")
        self.assertEqual(customer.street, "164/4 ม.6")
        self.assertEqual(customer.street2, "ตำบลสามพระยา")
        self.assertEqual(customer.city, "อำเภอชะอำ")
        self.assertEqual(customer.zip, "76120")
        self.assertEqual(customer.email, "buyer@example.com")
        self.assertEqual(self.order.partner_invoice_id, customer)

        delivery = self.order.partner_shipping_id
        self.assertNotEqual(delivery, customer)
        self.assertEqual(delivery.parent_id, customer)
        self.assertEqual(delivery.type, "delivery")
        self.assertEqual(delivery.name, "เค้ก")
        self.assertEqual(delivery.phone, "0644976555")
        self.assertEqual(delivery.street, "เลขที่ 63/50 หมู่ที่ 5")
        self.assertEqual(delivery.zip, "84320")
        self.assertIn("Tax invoice", wizard.result_message)

        # importing again reuses the same delivery contact
        wizard2 = self.env["shopee.buyer.address.import.wizard"].create({
            "file_name": "Order.toship.xlsx",
            "file_data": self._xlsx([self._row("260921KGN5CTYC", "A")]),
        })
        wizard2.action_import()
        self.assertEqual(self.order.partner_shipping_id, delivery)


MONEY_HEADERS = HEADERS + [
    "ราคาขายสุทธิ", "ค่าจัดส่งที่ชำระโดยผู้ซื้อ", "โค้ดส่วนลดชำระโดยผู้ขาย",
    "โค้ดส่วนลด", "ค่าคอมมิชชั่น", "ค่าบริการ", "Transaction Fee",
]


class TestShopeeBuyerPaymentImport(TestShopeeBuyerAddressImport):
    def _xlsx(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(MONEY_HEADERS)
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        return base64.b64encode(stream.getvalue())

    def _row(self, order_sn, product, line_total="5590.00"):
        return super()._row(order_sn, product) + [
            line_total, "200.00", "10.00", "DDX25HPHDL5SEP;MOGEN083",
            "1075.00", "478.00", "186.00",
        ]

    def test_import_applies_payment_details(self):
        wizard = self.env["shopee.buyer.address.import.wizard"].create({
            "file_name": "Order.shipping.xlsx",
            "file_data": self._xlsx([self._row("260921KGN5CTYC", "A")]),
        })
        wizard.action_import()
        voucher = self.order.order_line.filtered(lambda l: l.shopee_line_type == "voucher")
        self.assertIn("MOGEN083", voucher.name)
        self.assertNotIn("DDX25HPHDL5SEP", voucher.name)
        self.assertTrue(self.order.order_line.filtered(
            lambda l: l.shopee_line_type == "shipping"))
        # 5590 + 200 - 10 - (1075 + 478 + 186)
        self.assertEqual(self.order.shopee_net_income, 4041.0)
        self.assertIn("Payment details", wizard.result_message)
